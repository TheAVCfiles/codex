import datetime
import hashlib
import os
import re
from pathlib import Path

import pandas as pd
import requests
from bs4 import BeautifulSoup
from flask import Flask, jsonify, request
from flask_cors import CORS
from openpyxl import Workbook, load_workbook

app = Flask(__name__)
CORS(app)

BASE_DIR = Path(__file__).resolve().parent
INTAKE_DIR = BASE_DIR / "intake"
EXPORTS_DIR = BASE_DIR / "exports"
MASTER_XLSX = BASE_DIR / "case_master.xlsx"

for directory in (INTAKE_DIR, EXPORTS_DIR):
    directory.mkdir(parents=True, exist_ok=True)


def init_workbook() -> None:
    if MASTER_XLSX.exists():
        return

    wb = Workbook()
    sheets = [
        "Dashboard",
        "Intake_Log",
        "Raw_Entities",
        "Canonical_Entities",
        "Events",
        "Review_Queue",
    ]

    for index, title in enumerate(sheets):
        if index == 0:
            wb.active.title = title
        else:
            wb.create_sheet(title)

    headers = {
        "Intake_Log": [
            "doc_id",
            "filename",
            "source_type",
            "source_path",
            "import_date",
            "hash",
            "status",
        ],
        "Raw_Entities": [
            "raw_id",
            "doc_id",
            "raw_text",
            "type",
            "page",
            "confidence",
            "status",
        ],
        "Canonical_Entities": [
            "entity_id",
            "canonical_name",
            "type",
            "aliases",
            "doc_count",
            "status",
        ],
        "Review_Queue": [
            "item_id",
            "doc_id",
            "type",
            "content",
            "suggested_action",
            "status",
        ],
    }

    for sheet, cols in headers.items():
        wb[sheet].append(cols)

    wb.save(MASTER_XLSX)


def get_file_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def extract_signals(text: str) -> list[dict[str, str]]:
    patterns = {
        "Person": r"\b[A-Z][a-z]+ [A-Z][a-z]+\b",
        "Organization": r"\b[A-Z]{2,}\b|\b[A-Z][a-z]+ (?:Corp|Inc|LLC|Ltd|Bank|University)\b",
        "Date": r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4})|(?:(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]* \d{1,2},? \d{4})\b",
        "Money": r"\$\d+(?:,\d{3})*(?:\.\d{2})?\b",
        "Email": r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
    }

    found: list[dict[str, str]] = []
    for label, pattern in patterns.items():
        for match in sorted(set(re.findall(pattern, text))):
            found.append({"text": match, "type": label})
    return found


init_workbook()


@app.route("/api/intake/upload", methods=["POST"])
def upload_file():
    if "file" not in request.files:
        return jsonify({"error": "No file part"}), 400

    file = request.files["file"]
    content = file.read()

    if not content:
        return jsonify({"error": "Empty file"}), 400

    file_hash = get_file_hash(content)
    doc_id = f"DOC-{file_hash[:8].upper()}"
    filename = file.filename or f"{doc_id}.bin"
    save_path = INTAKE_DIR / filename

    save_path.write_bytes(content)

    wb = load_workbook(MASTER_XLSX)
    ws = wb["Intake_Log"]
    ws.append(
        [
            doc_id,
            filename,
            "upload",
            str(save_path),
            datetime.datetime.now().isoformat(),
            file_hash,
            "new",
        ]
    )
    wb.save(MASTER_XLSX)

    return jsonify({"doc_id": doc_id, "status": "imported"})


@app.route("/api/intake/url", methods=["POST"])
def pull_url():
    data = request.get_json(silent=True) or {}
    url = data.get("url")

    if not url:
        return jsonify({"error": "Missing url"}), 400

    try:
        response = requests.get(url, timeout=10)
        response.raise_for_status()
        content = response.text

        file_hash = get_file_hash(content.encode("utf-8"))
        doc_id = f"URL-{file_hash[:8].upper()}"

        soup = BeautifulSoup(content, "html.parser")
        title = soup.title.string.strip() if soup.title and soup.title.string else url
        text_content = soup.get_text("\n", strip=True)

        filename = f"{doc_id}.txt"
        save_path = INTAKE_DIR / filename
        save_path.write_text(text_content, encoding="utf-8")

        wb = load_workbook(MASTER_XLSX)
        ws = wb["Intake_Log"]
        ws.append(
            [
                doc_id,
                title,
                "url",
                url,
                datetime.datetime.now().isoformat(),
                file_hash,
                "new",
            ]
        )
        wb.save(MASTER_XLSX)

        return jsonify({"doc_id": doc_id, "status": "fetched"})
    except Exception as err:  # noqa: BLE001
        return jsonify({"error": str(err)}), 500


@app.route("/api/scan", methods=["POST"])
def scan_docs():
    wb = load_workbook(MASTER_XLSX)
    log_ws = wb["Intake_Log"]
    raw_ws = wb["Raw_Entities"]

    new_docs = [
        row
        for row in log_ws.iter_rows(min_row=2, values_only=True)
        if len(row) > 6 and row[6] == "new"
    ]

    results = []
    for doc in new_docs:
        doc_id, _, _, source_path, *_ = doc
        if not source_path or not os.path.exists(source_path):
            continue

        with open(source_path, "r", encoding="utf-8", errors="ignore") as handle:
            text = handle.read()

        entities = extract_signals(text)
        for ent in entities:
            raw_id = f"ENT-{hashlib.md5(f'{doc_id}:{ent['text']}:{ent['type']}'.encode()).hexdigest()[:8].upper()}"
            raw_ws.append([raw_id, doc_id, ent["text"], ent["type"], "1", 0.8, "review"])
            results.append({"doc_id": doc_id, "entity": ent["text"], "type": ent["type"]})

    for row in log_ws.iter_rows(min_row=2):
        if row[6].value == "new":
            row[6].value = "parsed"

    wb.save(MASTER_XLSX)
    return jsonify({"scanned": len(new_docs), "entities_found": len(results), "results": results})


@app.route("/api/data/<sheet>", methods=["GET"])
def get_sheet_data(sheet: str):
    if not MASTER_XLSX.exists():
        return jsonify([])

    try:
        df = pd.read_excel(MASTER_XLSX, sheet_name=sheet)
    except Exception as err:  # noqa: BLE001
        return jsonify({"error": str(err)}), 400

    records = df.where(pd.notna(df), None).to_dict(orient="records")
    return jsonify(records)


if __name__ == "__main__":
    print("Forensic Console Backend running on http://127.0.0.1:5000")
    app.run(port=5000, debug=True)
